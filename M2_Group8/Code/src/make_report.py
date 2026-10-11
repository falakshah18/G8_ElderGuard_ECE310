#!/usr/bin/env python3
"""
ElderGuard - report builder (M2)
ECE 310 Group 8

Generates the M2 report from Results/metrics.json so the numbers in the report
and the numbers in the results are always the same. Produces a .docx (and, if
LibreOffice is available, a .pdf next to it).

    python make_report.py --metrics ../../Results/metrics.json \
                          --out ../../Report/M2_Group8_Report.docx
"""

import argparse
import json
import os
import shutil
import subprocess
import sys

from docx import Document
from docx.shared import Pt

MODEL_ORDER = ["Threshold (amp-variance)", "Random Forest", "MLP"]
TEAM = [
    ("AU2440014", "Falak Shah", "B.Tech CSE"),
    ("AU2440228", "Hiranshee Doshi", "B.Tech CSE"),
    ("AU2440275", "Nikunj Kalariya", "B.Tech CSE"),
    ("AU2440229", "Aaryan Parmar", "B.Tech CSE"),
    ("AU2440171", "Priyanshi Verma", "B.Tech CSE"),
]


def _heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def _para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


def build(metrics_path, out_path):
    with open(metrics_path) as fh:
        r = json.load(fh)

    best = max(
        (m for m in MODEL_ORDER if m in r),
        key=lambda m: r[m]["accuracy"],
    )

    doc = Document()
    doc.add_heading("ElderGuard - M2: Technical Artifact", 0)
    _para(doc, "ECE 310 - Wireless Communications, Monsoon Semester 2026")
    _para(doc, "Milestone 2 - Technical Artifact | Project Category: "
               "Minimum Viable Product (MVP) | Group 8")
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Roll No.", "Name", "Program"
    for roll, name, prog in TEAM:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = roll, name, prog

    _heading(doc, "1. Abstract")
    _para(doc,
          "ElderGuard is a low-cost, privacy-preserving fall-detection system "
          "built on commodity Wi-Fi hardware. Instead of a wearable or a "
          "camera, it uses the Channel State Information (CSI) of an ordinary "
          "802.11n link to sense human movement. At M2 we implement and "
          "evaluate the system's first working core function - counting the "
          "number of people in a small room - which is the presence/occupancy "
          "foundation any fall detector needs.")
    _para(doc,
          f"Our best model, {best}, reached "
          f"{r[best]['accuracy'] * 100:.1f}% accuracy and "
          f"{r[best]['count_mae']:.3f} count-MAE over "
          f"{r.get('_n_windows', r[best]['n_windows'])} analysis windows "
          f"({r.get('_n_sessions', '?')} sessions).")

    _heading(doc, "2. Results")
    _para(doc, f"Dataset: {r.get('_dataset', 'n/a')}. Evaluation: "
               "leave-one-session-out (no window from a held-out session is "
               "seen during training).")
    t = doc.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    h = t.rows[0].cells
    h[0].text, h[1].text, h[2].text, h[3].text = (
        "Model", "Accuracy", "Macro-F1", "Count-MAE")
    for m in MODEL_ORDER:
        if m not in r:
            continue
        cells = t.add_row().cells
        cells[0].text = m
        cells[1].text = f"{r[m]['accuracy']:.4f}"
        cells[2].text = f"{r[m]['macro_f1']:.4f}"
        cells[3].text = f"{r[m]['count_mae']:.4f}"

    _heading(doc, "3. Figures")
    report_dir = os.path.dirname(os.path.abspath(out_path))
    results_dir = os.path.join(report_dir, os.pardir, "Results")
    for fig in ["model_comparison.png", "confusion_matrix_random_forest.png",
                "confusion_matrix_mlp.png", "feature_importance.png",
                "amp_variance_vs_people.png"]:
        fig_path = os.path.join(results_dir, fig)
        if os.path.exists(fig_path):
            doc.add_paragraph(fig.replace(".png", "").replace("_", " "))
            doc.add_picture(fig_path, width=Pt(430))

    _heading(doc, "4. Discussion and Limitations")
    _para(doc,
          "All three models beat the 3-class chance level of 33%. Accuracy is "
          "limited by the 1-vs-2-person boundary and by room-to-room "
          "variation, which is the cross-environment difficulty reported in "
          "the SOTA literature. The data is real, cited CSI from the "
          "RadioPoints dataset, not from our own testbed; capture on our "
          "ESP32-WROOM-32D + ESP8266 pair is the immediate M3 task.")

    os.makedirs(report_dir, exist_ok=True)
    doc.save(out_path)
    print(f"Wrote {out_path}")

    # Best-effort PDF conversion.
    pdf_path = os.path.splitext(out_path)[0] + ".pdf"
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf",
             "--outdir", report_dir, out_path],
            check=False,
        )
        print(f"PDF conversion attempted -> {pdf_path}")
    else:
        print("LibreOffice not found; export the .docx to PDF manually.")
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--metrics", default="../../Results/metrics.json")
    ap.add_argument("--out", default="../../Report/M2_Group8_Report.docx")
    a = ap.parse_args()
    if not os.path.exists(a.metrics):
        sys.exit(f"metrics.json not found at {a.metrics}")
    build(a.metrics, a.out)


if __name__ == "__main__":
    main()
